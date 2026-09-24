#!/usr/bin/env python3
"""Reproducible pre-release verification for Inventory Toolkit.

The default gate is lightweight and checks compilation, unit tests, logical
integrity, and every checked-in profile. ``--demo`` additionally executes all
three demo workflows in isolated subprocesses and validates their workbook
structure without modifying repository data.
"""

from __future__ import annotations

import argparse
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
PROFILES_ROOT = ROOT / "profiles"
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def _run(command: list[str]) -> None:
    print("  $", " ".join(command), flush=True)
    subprocess.run(command, cwd=ROOT, check=True)


def _quick_gate() -> None:
    _run([
        sys.executable,
        "-m",
        "compileall",
        "-q",
        "core",
        "cli",
        "engine",
        "gui",
        "tools",
        "tests",
    ])
    _run([sys.executable, "-m", "pytest", "-q"])
    _run([sys.executable, "tools/IntegrityCheck.py"])

    retirement_tool = ROOT / "tools" / "RetireLegacyCompatibility.py"
    if retirement_tool.exists():
        _run([sys.executable, str(retirement_tool.relative_to(ROOT))])

    _validate_profiles()


def _validate_profiles() -> None:
    os.chdir(ROOT)
    from core.configuration_manager import ConfigurationManager

    profiles = sorted(path.name for path in PROFILES_ROOT.iterdir() if path.is_dir())
    if not profiles:
        raise RuntimeError("No profiles were found under profiles/.")

    print("\nProfile validation:")
    for profile in profiles:
        config = ConfigurationManager(profile)
        config.get_catalog()
        config.get_family_config()
        config.get_network_config()
        config.get_stock_processing_config()
        config.get_cross_check_config()
        config.get_yoy_reports_config()
        print(f"  PASS  {profile}")


def _demo_gate(cross_system_override: str | None) -> None:
    print("\nDemo workflow validation:")
    for workflow in ("cross", "stock", "yoy"):
        command = [
            sys.executable,
            "tools/ReleaseCheck.py",
            "--_demo-workflow",
            workflow,
        ]
        if workflow == "cross" and cross_system_override:
            command.extend(["--cross-system", cross_system_override])
        _run(command)


def _run_internal_demo(workflow: str, cross_system_override: str | None) -> None:
    os.chdir(ROOT)
    if workflow == "cross":
        _run_cross_demo(cross_system_override)
    elif workflow == "stock":
        _run_stock_demo()
    elif workflow == "yoy":
        _run_yoy_demo()
    else:
        raise ValueError(f"Unknown demo workflow: {workflow}")


def _run_cross_demo(cross_system_override: str | None) -> None:
    from openpyxl import load_workbook

    from core.business_schema import (
        ARTICLE_COLUMN,
        COST_TOTAL_COLUMN,
        DIFFERENCE_COLUMN,
        FAMILY_COLUMN,
        PHYSICAL_COUNT_COLUMN,
        SALES_TOTAL_COLUMN,
        SYSTEM_STOCK_COLUMN,
    )
    from engine.inventory_cross_check.generator import run_cross_check

    demo = ROOT / "examples" / "demo"
    cross_system = (
        Path(cross_system_override).resolve()
        if cross_system_override
        else demo / "cross_check_system_stock.xls"
    )

    with tempfile.TemporaryDirectory(prefix="inventory_toolkit_cross_") as temp_dir:
        output = Path(temp_dir) / "cross_check.xlsx"
        result = run_cross_check(
            SimpleNamespace(
                cross_check_system=str(cross_system),
                cross_check_count=str(demo / "cross_check_physical_count.xlsx"),
                shared_cost=str(demo / "shared_cost_list.xlsx"),
                shared_sales=str(demo / "shared_sales_price_list.xlsx"),
                cross_check_out=str(output),
                cross_check_profile="demo",
                cross_check_consolidate=True,
                cross_check_partial=False,
                non_interactive=True,
            )
        )
        if not result:
            raise RuntimeError("Cross Check demo did not produce an output workbook.")
        workbook = load_workbook(result, data_only=False, read_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        expected_headers = [
            FAMILY_COLUMN,
            ARTICLE_COLUMN,
            SYSTEM_STOCK_COLUMN,
            PHYSICAL_COUNT_COLUMN,
            DIFFERENCE_COLUMN,
            COST_TOTAL_COLUMN,
            SALES_TOTAL_COLUMN,
        ]
        if headers != expected_headers:
            raise RuntimeError(f"Cross Check demo headers changed: {headers!r}")
        print(f"  PASS  Cross Check ({worksheet.max_row - 1} difference rows)")
        workbook.close()


def _run_stock_demo() -> None:
    from openpyxl import load_workbook

    from core.configuration_manager import ConfigurationManager
    from engine.stock_processing.generator import run_stock_processing

    demo = ROOT / "examples" / "demo"
    with tempfile.TemporaryDirectory(prefix="inventory_toolkit_stock_") as temp_dir:
        output = Path(temp_dir) / "stock.xlsx"
        result = run_stock_processing(
            SimpleNamespace(
                stock_processing_raw=str(demo / "stock_processing_raw_stock.xlsx"),
                shared_cost=str(demo / "shared_cost_list.xlsx"),
                shared_sales=str(demo / "shared_sales_price_list.xlsx"),
                stock_processing_out=str(output),
                stock_processing_profile="demo",
                non_interactive=True,
            )
        )
        if not result:
            raise RuntimeError("Stock Processing demo did not produce an output workbook.")
        stock_config = ConfigurationManager("demo").get_stock_output()
        workbook = load_workbook(result, data_only=False, read_only=True)
        expected_sheets = {
            stock_config["raw_data_sheet"],
            *(summary["sheet_name"] for summary in stock_config["summaries"]),
        }
        missing_sheets = expected_sheets.difference(workbook.sheetnames)
        if missing_sheets:
            raise RuntimeError(f"Stock Processing demo is missing sheets: {sorted(missing_sheets)}")
        raw_sheet = workbook[stock_config["raw_data_sheet"]]
        if raw_sheet.max_row <= 1:
            raise RuntimeError("Stock Processing raw-data sheet contains no data rows.")
        print(f"  PASS  Stock Processing ({raw_sheet.max_row - 1} data rows)")
        workbook.close()


def _run_yoy_demo() -> None:
    import pandas as pd
    from openpyxl import load_workbook

    from core.configuration_manager import ConfigurationManager
    from engine.yoy_reports.generator import generate_sales_report

    demo = ROOT / "examples" / "demo"
    yoy_config = ConfigurationManager("demo").get_yoy_reports_config()
    date_column = yoy_config["input"]["date_column"]
    history = pd.read_excel(
        demo / "yoy_sales_history.xlsx",
        sheet_name=0,
        usecols=[date_column],
    )
    dates = pd.to_datetime(history[date_column], errors="coerce").dropna()
    latest_year = int(dates.dt.year.max())
    current_dates = dates[dates.dt.year == latest_year]
    start_date = pd.Timestamp(year=latest_year, month=1, day=1)
    end_date = current_dates.max()

    with tempfile.TemporaryDirectory(prefix="inventory_toolkit_yoy_") as temp_dir:
        output = Path(temp_dir) / "yoy.xlsx"
        result = generate_sales_report(
            str(demo / "yoy_sales_history.xlsx"),
            str(output),
            start_date,
            end_date,
            yoy_config,
            yoy_config["input"]["grouping_column"],
            True,
            False,
            "demo",
            yoy_config["output"].get("include_sizes", False),
            True,
        )
        if not result:
            raise RuntimeError("YoY demo did not produce an output workbook.")
        workbook = load_workbook(result, data_only=False, read_only=True)
        if "Full Report" not in workbook.sheetnames:
            raise RuntimeError("YoY demo is missing the Full Report sheet.")
        if len(workbook.sheetnames) < 2:
            raise RuntimeError("YoY segmented demo did not create period sheets.")
        print(f"  PASS  YoY Reports ({len(workbook.sheetnames)} sheets)")
        workbook.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--demo",
        action="store_true",
        help="also execute all three example workflows end-to-end",
    )
    parser.add_argument(
        "--cross-system",
        metavar="PATH",
        help="override the demo Cross Check system-stock file (useful for converted .xlsx fixtures)",
    )
    parser.add_argument(
        "--_demo-workflow",
        choices=("cross", "stock", "yoy"),
        help=argparse.SUPPRESS,
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        if args._demo_workflow:
            _run_internal_demo(args._demo_workflow, args.cross_system)
            return 0

        print("Inventory Toolkit release verification")
        print(f"Repository: {ROOT}\n")
        _quick_gate()
        if args.demo:
            _demo_gate(args.cross_system)
    except (Exception, subprocess.CalledProcessError) as exc:
        print(f"\nRELEASE CHECK FAILED: {exc}", file=sys.stderr)
        return 1

    print("\nRELEASE CHECK PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
