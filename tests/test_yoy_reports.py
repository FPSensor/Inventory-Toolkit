import pandas as pd
import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from core.config_schemas import YoYReportsConfig
from engine.yoy_reports.generator import generate_sales_report
from engine.yoy_reports.metrics import configured_branches, resolve_metric_specs
from engine.yoy_reports.sheet_renderer import render_report_sheet


def _config(*, metrics=None, annual_comparison=True, include_sizes=False, groups=None):
    return {
        "version": 3,
        "input": {
            "date_column": "Fecha",
            "quantity_column": "Cantidad",
            "sales_column": "Monto",
            "grouping_column": "Familias",
            "item_column": "Articulo",
            "branch_column": "Base",
            "size_column": "Talle",
        },
        "output": {
            "default_path": "report.xlsx",
            "metrics": ["units", "sales"] if metrics is None else metrics,
            "annual_comparison": annual_comparison,
            "include_sizes": include_sizes,
        },
        "groups": groups if groups is not None else {"north": ["A", "B"]},
    }


def _frames():
    current = pd.DataFrame(
        {
            "Familias": ["Remeras", "Remeras", "Jeans"],
            "Articulo": ["001", "002", "130"],
            "Base": ["A", "B", "A"],
            "Cantidad": [2, 1, 3],
            "Monto": [200.50, 120.00, 450.25],
            "Talle": ["M", "L", "42"],
            "Fecha": pd.to_datetime(["2026-01-05", "2026-01-06", "2026-01-07"]),
        }
    )
    previous = current.copy()
    previous["Fecha"] = previous["Fecha"] - pd.DateOffset(years=1)
    previous["Cantidad"] = [1, 2, 2]
    previous["Monto"] = [90.00, 210.00, 300.00]
    return current, previous


def _sheet_values(worksheet):
    return [
        cell.value
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]


def test_yoy_metric_config_is_explicit_and_validated():
    config = YoYReportsConfig.model_validate(_config()).model_dump()
    specs = resolve_metric_specs(config)
    assert [(spec.key, spec.column) for spec in specs] == [
        ("units", "Cantidad"),
        ("sales", "Monto"),
    ]

    with pytest.raises(ValidationError):
        YoYReportsConfig.model_validate(_config(metrics=[]))
    with pytest.raises(ValidationError):
        YoYReportsConfig.model_validate(_config(metrics=["profit"]))


def test_yoy_rejects_empty_report_groups_before_rendering():
    with pytest.raises(ValueError, match="at least one non-empty report group"):
        configured_branches(_config(groups={}))
    with pytest.raises(ValueError, match="at least one non-empty report group"):
        configured_branches(_config(groups={"empty": []}))


def test_yoy_renders_configured_units_and_sales_metrics():
    current, previous = _frames()
    workbook = Workbook()
    worksheet = workbook.active
    config = _config(annual_comparison=False)

    render_report_sheet(
        worksheet,
        current,
        previous,
        pd.Timestamp("2026-01-01"),
        pd.Timestamp("2026-01-31"),
        pd.Timestamp("2025-01-01"),
        config,
        "Familias",
        include_sizes=False,
    )

    values = _sheet_values(worksheet)
    assert any(str(value).startswith("Units Sold from") for value in values)
    assert any(str(value).startswith("Sales Amount from") for value in values)
    assert not any(str(value).startswith("YoY ") for value in values)

    sales_cells = [
        cell
        for row in worksheet.iter_rows()
        for cell in row
        if cell.value in {200.50, 120.00, 450.25}
    ]
    assert sales_cells
    assert all(cell.number_format == "#,##0.00" for cell in sales_cells)


def test_yoy_annual_comparison_setting_controls_comparison_blocks():
    current, previous = _frames()
    workbook = Workbook()
    worksheet = workbook.active
    config = _config(annual_comparison=True)

    render_report_sheet(
        worksheet,
        current,
        previous,
        pd.Timestamp("2026-01-01"),
        pd.Timestamp("2026-01-31"),
        pd.Timestamp("2025-01-01"),
        config,
        "Familias",
    )

    values = [str(value) for value in _sheet_values(worksheet)]
    assert any(value.startswith("YoY Units Sold Comparison") for value in values)
    assert any(value.startswith("YoY Sales Amount Comparison") for value in values)


def test_generate_sales_report_uses_profile_include_sizes_default(monkeypatch, tmp_path):
    current, previous = _frames()
    config = _config(metrics=["units"], include_sizes=True)
    captured = {}

    def fake_process(*_args, **_kwargs):
        return current, previous, pd.Timestamp("2025-01-01")

    def fake_render(*args, **kwargs):
        captured["include_sizes"] = args[9]
        return str(tmp_path / "report.xlsx")

    monkeypatch.setattr("engine.yoy_reports.generator.process_sales_data", fake_process)
    monkeypatch.setattr("engine.yoy_reports.generator.render_yoy_sales_excel", fake_render)

    generate_sales_report(
        "input.xlsx",
        str(tmp_path / "report.xlsx"),
        pd.Timestamp("2026-01-01"),
        pd.Timestamp("2026-01-31"),
        config,
        "Familias",
        False,
        True,
        "demo",
    )

    assert captured["include_sizes"] is True
