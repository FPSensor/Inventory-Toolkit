"""Exercise data-owned businesses through the shared setup service and real engines."""
from copy import deepcopy
from itertools import product
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
import pandas as pd
import pytest

from core import paths
from core.business_schema import DIFFERENCE_COLUMN, COST_TOTAL_COLUMN, SALES_TOTAL_COLUMN
from core.configuration_manager import ConfigurationManager
from core.profile_creation import create_profile
from core.setup_catalog import get_value
from core.setup_service import SetupSession, SetupState
from engine.stock_processing.generator import run_stock_processing
from engine.inventory_cross_check.generator import run_cross_check
from engine.yoy_reports.generator import generate_sales_report
from engine.yoy_reports.data_processor import process_sales_data

CASES = json.loads((Path(__file__).parents[1] / 'examples/businesses/scenarios.json').read_text())


@pytest.fixture(params=CASES, ids=lambda case: case['name'])
def business(request, tmp_path, monkeypatch):
    case = request.param
    monkeypatch.setattr(paths, 'PROFILES_ROOT', tmp_path / 'profiles')
    profile = create_profile(case['name'], 'Synthetic verification only', ['stock_processing', 'cross_check', 'yoy_reports'])
    for module in ('stock_processing', 'cross_check', 'yoy_reports'):
        session = SetupSession(profile, module, include_advanced=True)
        visited = set()
        while pending := [step for step in session.steps if step.id not in visited]:
            step = pending[0]
            desired = case['configuration'].get(step.logical, {})
            values = {}
            for field in step.fields:
                try:
                    value = get_value(desired, field.path)
                except (KeyError, TypeError):
                    value = get_value(session.data(step), field.path)
                values['.'.join(field.path)] = deepcopy(value)
            session.save_step(step.id, values)
            visited.add(step.id)
        session.validate_required(include_sales=module == 'yoy_reports', include_variants=module == 'yoy_reports')
        assert session.state() == SetupState.COMPLETE
    files = {}
    for role, rows in case['workbooks'].items():
        path = tmp_path / f'{role}.xlsx'
        pd.DataFrame(rows).to_excel(path, index=False)
        files[role] = str(path)
    return case, profile, files, tmp_path


def test_business_stock_quantities_categories_and_valuations(business):
    case, profile, files, root = business
    output = root / 'stock-report.xlsx'
    args = SimpleNamespace(stock_processing_raw=files['stock'], shared_cost=files['prices'],
        shared_sales=files['prices'], stock_processing_out=str(output), stock_processing_profile=profile,
        non_interactive=True)
    assert run_stock_processing(args) == str(output)
    frame = pd.read_excel(output, sheet_name='Stock', dtype={'SKU': str})
    assert frame['SKU'].tolist() == [row['Stock code'] for row in case['workbooks']['stock']]
    assert frame['Category'].tolist() == [row['Section'] for row in case['workbooks']['stock']]
    for location in case['configuration']['general/network']['active']:
        assert frame[location].sum() == pytest.approx(case['expected']['stock_per_location'])
        # Valuation roles use bundled vocabulary, never inferred from display labels.
        from core.business_schema import COST_COLUMN, SALES_VALUE_LABEL
        assert frame[f'{location}.{COST_COLUMN}'].sum() == pytest.approx(case['expected']['cost_per_location'])
        assert frame[f'{location}.{SALES_VALUE_LABEL}'].sum() == pytest.approx(case['expected']['sales_per_location'])


def test_business_cross_options_preserve_reconciliation(business):
    case, profile, files, root = business
    for consolidate, partial in product((False, True), repeat=2):
        output = root / f'cross-{consolidate}-{partial}.xlsx'
        args = SimpleNamespace(cross_check_system=files['system'], cross_check_count=files['count'],
            shared_cost=files['prices'], shared_sales=files['prices'], cross_check_out=str(output),
            cross_check_profile=profile, cross_check_consolidate=consolidate,
            cross_check_partial=partial, non_interactive=True)
        assert run_cross_check(args) == str(output)
        frame = pd.read_excel(output)
        assert sorted(frame[DIFFERENCE_COLUMN]) == sorted(case['expected']['differences'])
        assert frame[COST_TOTAL_COLUMN].sum() == pytest.approx(case['expected']['difference_cost'])
        assert frame[SALES_TOTAL_COLUMN].sum() == pytest.approx(case['expected']['difference_sale'])


def test_business_yoy_option_matrix(business):
    case, profile, files, root = business
    config = ConfigurationManager(profile, module='yoy_reports').get_yoy_reports_config()
    start, end = pd.Timestamp('2026-01-01'), pd.Timestamp('2026-02-28')
    current, previous, _ = process_sales_data(files['sales'], start, end, config)
    assert current['Sold quantity'].sum() == pytest.approx(case['expected']['current_units_per_month'] * 2)
    assert previous['Sold quantity'].sum() == pytest.approx(case['expected']['previous_units_per_month'] * 2)
    assert current['Line total'].sum() == pytest.approx(case['expected']['current_sales_per_month'] * 2)
    for segmented, annual, variants, revenue, by_item, source_categories in product((False, True), repeat=6):
        selected = deepcopy(config)
        selected['output'].update(annual_comparison=annual, metrics=['units','sales'] if revenue else ['units'])
        output = root / f'yoy-{segmented}-{annual}-{variants}-{revenue}-{by_item}-{source_categories}.xlsx'
        grouping = selected['input']['item_column' if by_item else 'grouping_column']
        generate_sales_report(files['sales'], str(output), start, end, selected, grouping,
            segmented, source_categories, profile, yoy_include_variants=variants, non_interactive=True)
        workbook = load_workbook(output, data_only=False)
        assert len(workbook.sheetnames) == (3 if segmented else 1)
        # Every generated sheet has a real metric block, and variants never create sheets.
        for sheet in workbook:
            values = [cell.value for row in sheet for cell in row if cell.value is not None]
            title_rows = [row[0].row for row in sheet if str(row[0].value).startswith(selected['output']['quantity_label'] + ' from ')]
            assert len(title_rows) == 1
            first = title_rows[0] + 2
            last = next(row for row in range(first, sheet.max_row + 1) if sheet.cell(row, 1).value == 'Totals')
            locations = case['configuration']['general/network']['active']
            leaf_sum = sum(sheet.cell(row, col).value for row in range(first, last)
                for col in range(2, len(locations) + 2)
                if isinstance(sheet.cell(row, col).value, (int, float)))
            months = 2 if not segmented or sheet.title == 'Full Report' else 1
            assert leaf_sum == pytest.approx(case['expected']['current_units_per_month'] * months)
            if revenue:
                assert any(str(value).startswith('Sales Amount from ') for value in values)
            else:
                assert not any(str(value).startswith('Sales Amount from ') for value in values)
        workbook.close()
