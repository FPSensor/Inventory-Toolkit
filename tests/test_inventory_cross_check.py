import pytest
from engine.shared.families import build_family_rules, assign_family
from engine.inventory_cross_check.data_processor import calculate_difference, normalize_article

def test_build_family_rules():
    families = {
        "Remeras": ["001", "002"],
        "Buzos": ["0085", "185"]
    }
    rules = build_family_rules(families)
    # Longitud máxima primero
    assert len(rules[0][0]) == 4
    assert rules[0][0] == "0085"
    # Longitud mínima al final (3 caracteres)
    assert len(rules[-1][0]) == 3

def test_assign_family():
    rules = [("0085", "Buzos"), ("001", "Remeras")]
    assert assign_family("0085-123", rules) == "Buzos"
    assert assign_family("00100-XYZ", rules) == "Remeras"
    assert assign_family("99999-ABC", rules) == "Other"
    assert assign_family("REVISAR | 123", rules) == "REVISAR"

def test_calculate_difference():
    # Stock 10, Conteo 5 -> Faltan 5
    assert calculate_difference(10, 5) == -5
    # System stock -2, physical count 5 -> difference is 5; negative system stock is ignored.
    assert calculate_difference(-2, 5) == 5
    # Stock 0, Conteo 10 -> Sobran 10
    assert calculate_difference(0, 10) == 10


def test_normalize_article_variable_length_longest_prefix():
    # Scanner data may append size/color to variable-length article codes.
    # The master stock is the source of truth; no fixed substring length is valid.
    master_base = [
        "11111",
        "11111-2",
        "11111-261",
        "LONG-ARTICLE-12345",
    ]
    master_set = {article.upper() for article in master_base}

    assert normalize_article("11111-261XXLH1", master_base, master_set) == "11111-261"
    assert normalize_article("LONG-ARTICLE-12345XLRED", master_base, master_set) == "LONG-ARTICLE-12345"


def test_normalize_article_exact_match_and_review_fallback():
    master_base = ["11111-261", "ABC-123"]
    master_set = {article.upper() for article in master_base}

    # Exact matches are normalized to uppercase, matching system-stock normalization.
    assert normalize_article("abc-123", master_base, master_set) == "ABC-123"

    # Unknown readings must never be guessed, truncated, or discarded.
    raw_unknown = "UNKNOWN-XXLH1"
    assert normalize_article(raw_unknown, master_base, master_set) == f"REVISAR | {raw_unknown}"


def test_normalize_article_ignores_empty_master_entries():
    master_base = ["", "   ", "ABC", "ABC-123"]
    master_set = {str(article).upper().strip() for article in master_base}

    # A zero-length master article must never win prefix matching. Unknown
    # scanner readings must remain visible for manual review.
    raw_unknown = "UNKNOWN-XXLH1"
    assert normalize_article(raw_unknown, master_base, master_set) == f"REVISAR | {raw_unknown}"

    # Existing invariants still hold in the presence of invalid blank masters.
    assert normalize_article("abc", master_base, master_set) == "ABC"
    assert normalize_article("ABC-123XLRED", master_base, master_set) == "ABC-123"


def test_cross_check_renderer_owns_header_style(tmp_path):
    """Cross Check headers must not inherit version-dependent Pandas styling."""
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    from engine.inventory_cross_check.excel_renderer import apply_excel_formatting

    output = tmp_path / "cross.xlsx"
    pd.DataFrame(
        [{
            "Familias": "Demo",
            "Artículo": "ABC-1",
            "Stock Sistema": 1,
            "Conteo Físico": 2,
            "Diferencia": 1,
            "CTOTAL": 10.0,
            "VTOTAL": 20.0,
        }]
    ).to_excel(output, index=False)

    # Simulate a different Pandas/OpenPyXL runtime contributing a different
    # header style before Inventory Toolkit applies its own presentation.
    workbook = load_workbook(output)
    for cell in workbook.active[1]:
        cell.font = Font(name="Arial", size=14, italic=True, color="FF0000")
        cell.fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    workbook.save(output)
    workbook.close()

    apply_excel_formatting(output, interactive=False)

    workbook = load_workbook(output)
    try:
        for cell in workbook.active[1]:
            assert cell.font.name == "Calibri"
            assert cell.font.sz == 11
            assert cell.font.bold is True
            assert cell.font.italic is False
            assert cell.font.color.type == "theme"
            assert cell.font.color.theme == 1
            assert cell.fill.fill_type is None
            assert cell.alignment.horizontal == "center"
            assert cell.alignment.vertical == "center"
            assert cell.number_format == "General"
            assert cell.protection.locked is True
            assert cell.protection.hidden is False
            assert cell.border.left.style == "thin"
            assert cell.border.right.style == "thin"
            assert cell.border.top.style == "thin"
            assert cell.border.bottom.style == "thin"
    finally:
        workbook.close()
