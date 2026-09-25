"""Worksheet orchestration for Year-over-Year sales reports."""

from openpyxl.utils import get_column_letter

from core.business_schema import SIZE_COLUMN
from engine.yoy_reports.comparison_renderer import render_yoy_comparison_blocks
from engine.yoy_reports.current_sales_renderer import render_current_sales_block
from engine.yoy_reports.metrics import configured_branches, resolve_metric_specs


def render_report_sheet(
    worksheet,
    current_frame,
    previous_frame,
    start_date,
    end_date,
    previous_start,
    yoy_config,
    grouping_column,
    include_sizes=False,
):
    input_config = yoy_config["input"]
    output_config = yoy_config["output"]
    branch_column = input_config["branch_column"]
    size_column = input_config.get("size_column", SIZE_COLUMN)
    groups = yoy_config["groups"]
    all_branches = configured_branches(yoy_config)
    metric_specs = resolve_metric_specs(yoy_config)
    annual_comparison = output_config.get("annual_comparison", True)

    current_row = 1
    for metric in metric_specs:
        current_row = render_current_sales_block(
            worksheet,
            current_frame,
            start_date,
            end_date,
            grouping_column,
            branch_column,
            metric.column,
            size_column,
            all_branches,
            include_sizes,
            start_row=current_row,
            metric_label=metric.label,
            number_format=metric.number_format,
        )
        if annual_comparison:
            current_row = render_yoy_comparison_blocks(
                worksheet,
                current_frame,
                previous_frame,
                start_date,
                end_date,
                previous_start,
                grouping_column,
                branch_column,
                metric.column,
                groups,
                current_row,
                metric_label=metric.label,
                number_format=metric.number_format,
            )

    for column_index in range(1, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = 15
    worksheet.column_dimensions["A"].width = 30
