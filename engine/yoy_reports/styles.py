"""Shared cell styling primitives for YoY workbooks."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def apply_style(cell, is_header=False, is_total=False, num_format=None, indent=0):
    cell.alignment = Alignment(
        horizontal="center" if indent == 0 else "left",
        vertical="center",
        indent=indent,
    )
    cell.border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    if is_header:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            start_color="D3D3D3",
            end_color="D3D3D3",
            fill_type="solid",
        )
    if is_total:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            start_color="F0F0F0",
            end_color="F0F0F0",
            fill_type="solid",
        )
    if num_format:
        cell.number_format = num_format
