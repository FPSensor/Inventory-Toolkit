"""Render branch/group Year-over-Year comparison blocks."""

import pandas as pd
from openpyxl.utils import get_column_letter

from engine.yoy_reports.styles import apply_style


def render_yoy_comparison_blocks(
    worksheet,
    current_frame,
    previous_frame,
    start_date,
    end_date,
    previous_start,
    grouping_column,
    branch_column,
    quantity_column,
    groups,
    current_row,
):
    """Render all configured YoY comparison groups and return the next row."""
    for group_name, branches in groups.items():
        if not branches:
            continue

        # A Global block only exists when a group contains multiple branches.
        has_global = len(branches) > 1
        block_count = len(branches) + (1 if has_global else 0)
        group_width = 1 + block_count * 3

        worksheet.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=group_width,
        )
        apply_style(
            worksheet.cell(
                row=current_row,
                column=1,
                value=f'YoY Comparison - {group_name.replace("_", " ").title()}',
            ),
            is_header=True,
        )
        current_row += 1

        column_index = 2
        for branch in branches:
            worksheet.merge_cells(
                start_row=current_row,
                start_column=column_index,
                end_row=current_row,
                end_column=column_index + 1,
            )
            apply_style(
                worksheet.cell(row=current_row, column=column_index, value=branch),
                is_header=True,
            )
            apply_style(
                worksheet.cell(row=current_row, column=column_index + 1),
                is_header=True,
            )
            apply_style(
                worksheet.cell(row=current_row, column=column_index + 2),
                is_header=True,
            )
            column_index += 3

        if has_global:
            worksheet.merge_cells(
                start_row=current_row,
                start_column=column_index,
                end_row=current_row,
                end_column=column_index + 1,
            )
            apply_style(
                worksheet.cell(row=current_row, column=column_index, value="Global"),
                is_header=True,
            )
            apply_style(
                worksheet.cell(row=current_row, column=column_index + 1),
                is_header=True,
            )
            apply_style(
                worksheet.cell(row=current_row, column=column_index + 2),
                is_header=True,
            )

        current_row += 1

        apply_style(
            worksheet.cell(row=current_row, column=1, value=grouping_column),
            is_header=True,
        )
        column_index = 2
        previous_year = previous_start.year
        current_year = end_date.year if pd.notna(end_date) else start_date.year + 1

        for _branch in branches:
            apply_style(
                worksheet.cell(
                    row=current_row,
                    column=column_index,
                    value=previous_year,
                ),
                is_header=True,
            )
            apply_style(
                worksheet.cell(
                    row=current_row,
                    column=column_index + 1,
                    value=current_year,
                ),
                is_header=True,
            )
            apply_style(
                worksheet.cell(row=current_row, column=column_index + 2, value="%"),
                is_header=True,
            )
            column_index += 3

        if has_global:
            apply_style(
                worksheet.cell(
                    row=current_row,
                    column=column_index,
                    value=previous_year,
                ),
                is_header=True,
            )
            apply_style(
                worksheet.cell(
                    row=current_row,
                    column=column_index + 1,
                    value=current_year,
                ),
                is_header=True,
            )
            apply_style(
                worksheet.cell(row=current_row, column=column_index + 2, value="%"),
                is_header=True,
            )

        previous_pivot = pd.pivot_table(
            previous_frame[previous_frame[branch_column].isin(branches)],
            values=quantity_column,
            index=grouping_column,
            columns=branch_column,
            aggfunc="sum",
            fill_value=0,
        )
        current_pivot = pd.pivot_table(
            current_frame[current_frame[branch_column].isin(branches)],
            values=quantity_column,
            index=grouping_column,
            columns=branch_column,
            aggfunc="sum",
            fill_value=0,
        )

        all_items = list(
            dict.fromkeys(list(previous_pivot.index) + list(current_pivot.index))
        )
        all_items.sort()

        for branch in branches:
            if branch not in previous_pivot.columns:
                previous_pivot[branch] = 0
            if branch not in current_pivot.columns:
                current_pivot[branch] = 0

        previous_pivot = previous_pivot.reindex(all_items, fill_value=0)
        current_pivot = current_pivot.reindex(all_items, fill_value=0)

        current_row += 1
        group_start_row = current_row
        for item in all_items:
            apply_style(worksheet.cell(row=current_row, column=1, value=item))
            column_index = 2
            previous_cells = []
            current_cells = []

            for branch in branches:
                previous_value = previous_pivot.loc[item, branch]
                current_value = current_pivot.loc[item, branch]
                apply_style(
                    worksheet.cell(
                        row=current_row,
                        column=column_index,
                        value=previous_value if previous_value != 0 else None,
                    ),
                    num_format="#,##0",
                )
                apply_style(
                    worksheet.cell(
                        row=current_row,
                        column=column_index + 1,
                        value=current_value if current_value != 0 else None,
                    ),
                    num_format="#,##0",
                )

                previous_letter = get_column_letter(column_index)
                current_letter = get_column_letter(column_index + 1)
                previous_cells.append(f"{previous_letter}{current_row}")
                current_cells.append(f"{current_letter}{current_row}")

                apply_style(
                    worksheet.cell(
                        row=current_row,
                        column=column_index + 2,
                        value=(
                            f'=IF({previous_letter}{current_row}=0, "N/A", '
                            f'({current_letter}{current_row}-{previous_letter}{current_row})/'
                            f'{previous_letter}{current_row})'
                        ),
                    ),
                    num_format="0.00%",
                )
                column_index += 3

            if has_global:
                global_previous_letter = get_column_letter(column_index)
                global_current_letter = get_column_letter(column_index + 1)
                apply_style(
                    worksheet.cell(
                        row=current_row,
                        column=column_index,
                        value="=" + "+".join(previous_cells),
                    ),
                    num_format="#,##0",
                )
                apply_style(
                    worksheet.cell(
                        row=current_row,
                        column=column_index + 1,
                        value="=" + "+".join(current_cells),
                    ),
                    num_format="#,##0",
                )
                apply_style(
                    worksheet.cell(
                        row=current_row,
                        column=column_index + 2,
                        value=(
                            f'=IF({global_previous_letter}{current_row}=0, "N/A", '
                            f'({global_current_letter}{current_row}-{global_previous_letter}{current_row})/'
                            f'{global_previous_letter}{current_row})'
                        ),
                    ),
                    num_format="0.00%",
                )

            current_row += 1

        group_end_row = current_row - 1
        if group_end_row < group_start_row:
            group_end_row = group_start_row

        apply_style(
            worksheet.cell(row=current_row, column=1, value="Totals"),
            is_total=True,
        )
        column_index = 2
        for _branch in branches:
            previous_letter = get_column_letter(column_index)
            current_letter = get_column_letter(column_index + 1)
            apply_style(
                worksheet.cell(
                    row=current_row,
                    column=column_index,
                    value=(
                        f"=SUM(${previous_letter}{group_start_row}:"
                        f"${previous_letter}{group_end_row})"
                    ),
                ),
                is_total=True,
                num_format="#,##0",
            )
            apply_style(
                worksheet.cell(
                    row=current_row,
                    column=column_index + 1,
                    value=(
                        f"=SUM(${current_letter}{group_start_row}:"
                        f"${current_letter}{group_end_row})"
                    ),
                ),
                is_total=True,
                num_format="#,##0",
            )
            apply_style(
                worksheet.cell(
                    row=current_row,
                    column=column_index + 2,
                    value=(
                        f'=IF({previous_letter}{current_row}=0, "N/A", '
                        f'({current_letter}{current_row}-{previous_letter}{current_row})/'
                        f'{previous_letter}{current_row})'
                    ),
                ),
                is_total=True,
                num_format="0.00%",
            )
            column_index += 3

        if has_global:
            global_previous_letter = get_column_letter(column_index)
            global_current_letter = get_column_letter(column_index + 1)
            apply_style(
                worksheet.cell(
                    row=current_row,
                    column=column_index,
                    value=(
                        f"=SUM(${global_previous_letter}{group_start_row}:"
                        f"${global_previous_letter}{group_end_row})"
                    ),
                ),
                is_total=True,
                num_format="#,##0",
            )
            apply_style(
                worksheet.cell(
                    row=current_row,
                    column=column_index + 1,
                    value=(
                        f"=SUM(${global_current_letter}{group_start_row}:"
                        f"${global_current_letter}{group_end_row})"
                    ),
                ),
                is_total=True,
                num_format="#,##0",
            )
            apply_style(
                worksheet.cell(
                    row=current_row,
                    column=column_index + 2,
                    value=(
                        f'=IF({global_previous_letter}{current_row}=0, "N/A", '
                        f'({global_current_letter}{current_row}-{global_previous_letter}{current_row})/'
                        f'{global_previous_letter}{current_row})'
                    ),
                ),
                is_total=True,
                num_format="0.00%",
            )

        current_row += 3

    return current_row
