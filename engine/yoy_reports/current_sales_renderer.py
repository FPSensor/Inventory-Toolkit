"""Render the current-period sales block of a YoY worksheet."""

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from engine.yoy_reports.styles import apply_style


def render_current_sales_block(
    worksheet,
    current_frame,
    start_date,
    end_date,
    grouping_column,
    branch_column,
    quantity_column,
    size_column,
    all_branches,
    include_sizes=False,
):
    """Render current-period totals and return the first row after the block."""
    main_block_width = len(all_branches) + 2

    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=main_block_width,
    )
    title_cell = worksheet["A1"]
    if pd.isna(start_date) or pd.isna(end_date):
        title_cell.value = "Sales (No Data)"
    else:
        title_cell.value = (
            f'Sales from {start_date.strftime("%B %d, %Y")} '
            f'to {end_date.strftime("%B %d, %Y")}'
        )
    apply_style(title_cell, is_header=True)

    main_headers = [grouping_column] + all_branches + ["Totals"]
    for column_index, header in enumerate(main_headers, 1):
        apply_style(
            worksheet.cell(row=2, column=column_index, value=header),
            is_header=True,
        )

    current_row = 3
    main_start_row = current_row

    if include_sizes and size_column in current_frame.columns:
        worksheet.sheet_properties.outlinePr.summaryBelow = False
        clean_current_frame = current_frame.copy()
        clean_current_frame[size_column] = (
            clean_current_frame[size_column].fillna("-").astype(str)
        )

        pivot = pd.pivot_table(
            clean_current_frame,
            values=quantity_column,
            index=[grouping_column, size_column],
            columns=branch_column,
            aggfunc="sum",
            fill_value=0,
        )
        for branch in all_branches:
            if branch not in pivot.columns:
                pivot[branch] = 0
        pivot = pivot[all_branches]
        families = pivot.index.get_level_values(0).unique().sort_values()

        for family in families:
            family_data = pivot.loc[family]
            family_start_row = current_row
            apply_style(
                worksheet.cell(row=current_row, column=1, value=family),
                is_header=False,
            )
            worksheet.cell(row=current_row, column=1).font = Font(bold=True)
            current_row += 1

            for size_value, row in family_data.iterrows():
                # "Talle" is part of the workbook contract, not an internal
                # identifier; formulas below intentionally match it.
                apply_style(
                    worksheet.cell(
                        row=current_row,
                        column=1,
                        value=f"  Talle: {size_value}",
                    ),
                    indent=1,
                )
                for index, branch in enumerate(all_branches):
                    value = row[branch]
                    apply_style(
                        worksheet.cell(
                            row=current_row,
                            column=2 + index,
                            value=value if value != 0 else None,
                        ),
                        num_format="#,##0",
                    )

                branch_letters = [
                    get_column_letter(2 + index)
                    for index in range(len(all_branches))
                ]
                apply_style(
                    worksheet.cell(
                        row=current_row,
                        column=main_block_width,
                        value="="
                        + "+".join(
                            f"{column_letter}{current_row}"
                            for column_letter in branch_letters
                        ),
                    ),
                    num_format="#,##0",
                )

                worksheet.row_dimensions[current_row].outlineLevel = 1
                worksheet.row_dimensions[current_row].hidden = True
                current_row += 1

            family_end_row = current_row - 1
            for index in range(len(all_branches) + 1):
                column_letter = get_column_letter(2 + index)
                apply_style(
                    worksheet.cell(
                        row=family_start_row,
                        column=2 + index,
                        value=(
                            f"=SUM(${column_letter}{family_start_row + 1}:"
                            f"${column_letter}{family_end_row})"
                        ),
                    ),
                    is_total=True,
                    num_format="#,##0",
                )
    else:
        current_pivot = pd.pivot_table(
            current_frame,
            values=quantity_column,
            index=grouping_column,
            columns=branch_column,
            aggfunc="sum",
            fill_value=0,
        )
        for branch in all_branches:
            if branch not in current_pivot.columns:
                current_pivot[branch] = 0
        current_pivot = current_pivot[all_branches].sort_index()

        for item, row in current_pivot.iterrows():
            apply_style(worksheet.cell(row=current_row, column=1, value=item))
            for index, branch in enumerate(all_branches):
                value = row[branch]
                apply_style(
                    worksheet.cell(
                        row=current_row,
                        column=2 + index,
                        value=value if value != 0 else None,
                    ),
                    num_format="#,##0",
                )

            branch_letters = [
                get_column_letter(2 + index)
                for index in range(len(all_branches))
            ]
            apply_style(
                worksheet.cell(
                    row=current_row,
                    column=main_block_width,
                    value="="
                    + "+".join(
                        f"{column_letter}{current_row}"
                        for column_letter in branch_letters
                    ),
                ),
                num_format="#,##0",
                is_total=True,
            )
            current_row += 1

    main_end_row = current_row - 1
    if main_end_row < main_start_row:
        main_end_row = main_start_row

    apply_style(
        worksheet.cell(row=current_row, column=1, value="Totals"),
        is_total=True,
    )
    for index in range(len(all_branches) + 1):
        column_letter = get_column_letter(2 + index)
        if include_sizes and size_column in current_frame.columns:
            formula = (
                f'=SUMIF($A${main_start_row}:$A${main_end_row}, "<>  Talle*", '
                f'${column_letter}${main_start_row}:${column_letter}${main_end_row})'
            )
        else:
            formula = (
                f"=SUM(${column_letter}${main_start_row}:"
                f"${column_letter}${main_end_row})"
            )
        apply_style(
            worksheet.cell(row=current_row, column=2 + index, value=formula),
            is_total=True,
            num_format="#,##0",
        )

    return current_row + 3
